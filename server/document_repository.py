
from pathlib import Path

SUPPORTED_EXTENSIONS = {".md", ".txt", ".pdf", ".ppt", ".pptx", ".docx", ".xlsx"}


class DocumentRepository:
    def __init__(
        self,
        docs_dir: Path,
        supported_extensions: set[str] | None = None,
    ):
        self.docs_dir = docs_dir.resolve()
        self.supported_extensions = (
            SUPPORTED_EXTENSIONS
            if supported_extensions is None
            else supported_extensions
        )

    def get_documents(self) -> list[Path]:
        self.docs_dir.mkdir(parents=True, exist_ok=True)

        return sorted(
            (
                p for p in self.docs_dir.rglob("*")
                if p.is_file()
                and p.suffix.lower() in self.supported_extensions
            ),
            key=lambda p: p.relative_to(self.docs_dir).as_posix(),
        )

    def read_document(self, path: Path) -> str:
        resolved = path.resolve()

        if not resolved.is_relative_to(self.docs_dir):
            raise ValueError(
                "Document path is outside the configured document directory"
            )

        extension = resolved.suffix.lower()

        if extension == ".pdf":
            try:
                import pymupdf
            except ImportError as exc:
                raise RuntimeError(
                    "PDF support requires PyMuPDF"
                ) from exc

            with pymupdf.open(resolved) as pdf:
                return "\n".join(
                    page.get_text() for page in pdf
                )

        if extension == ".pptx":
            try:
                from pptx import Presentation
            except ImportError as exc:
                raise RuntimeError(
                    "PPTX support requires python-pptx"
                ) from exc

            presentation = Presentation(str(resolved))
            slides_text = []

            for slide_number, slide in enumerate(
                presentation.slides, start=1
            ):
                texts = [
                    shape.text.strip()
                    for shape in slide.shapes
                    if shape.has_text_frame and shape.text.strip()
                ]

                if texts:
                    slides_text.append(
                        f"[슬라이드 {slide_number}]\n"
                        + "\n".join(texts)
                    )

            return "\n\n".join(slides_text)

        if extension == ".ppt":
            raise RuntimeError(
                "Legacy .ppt support is not implemented. "
                "Convert the file to .pptx first."
            )

        if extension == ".docx":
            try:
                from docx import Document
            except ImportError as exc:
                raise RuntimeError(
                    "DOCX support requires python-docx"
                ) from exc

            document = Document(str(resolved))
            contents = []

            contents.extend(
                paragraph.text
                for paragraph in document.paragraphs
                if paragraph.text.strip()
            )

            # 표
            for table in document.tables:
                for row in table.rows:
                    cells = [
                        cell.text.strip()
                        for cell in row.cells
                    ]
                    if any(cells):
                        contents.append(" | ".join(cells))

            return "\n".join(contents)

        if extension == ".xlsx":
            try:
                from openpyxl import load_workbook
            except ImportError as exc:
                raise RuntimeError(
                    "XLSX support requires openpyxl"
                ) from exc

            workbook = load_workbook(
                filename=resolved,
                read_only=True,
                data_only=True,
            )

            try:
                sheets_text = []

                for sheet in workbook.worksheets:
                    rows_text = []

                    for row in sheet.iter_rows(values_only=True):
                        values = [
                            str(value) if value is not None else ""
                            for value in row
                        ]

                        # 빈 행 제외
                        if any(value.strip() for value in values):
                            rows_text.append(" | ".join(values))

                    if rows_text:
                        sheets_text.append(
                            f"[시트: {sheet.title}]\n"
                            + "\n".join(rows_text)
                        )

                return "\n\n".join(sheets_text)
            finally:
                workbook.close()

        return resolved.read_text(encoding="utf-8")

    def relative_name(self, path: Path) -> str:
        return path.relative_to(self.docs_dir).as_posix()
